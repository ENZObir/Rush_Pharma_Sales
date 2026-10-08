"""Ouvre le template et remplace uniquement le contenu de ses tableaux nommés.

Le reste du classeur (feuille Outil, formules, listes déroulantes, mises en forme) n'est jamais touché.
openpyxl ne conserve pas les graphiques d'un classeur existant : le template n'en contient pas.
"""

import datetime as dt
import math
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.table import TableColumn


def _cell_value(v):
    if v is None or v is pd.NA or (isinstance(v, float) and math.isnan(v)):
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    if isinstance(v, pd.Period):
        return str(v)
    if hasattr(v, "item"):  # scalaires numpy
        return v.item()
    return v


class TemplateWriter:
    def __init__(self, template: Path, output: Path):
        self.template = template
        self.output = output
        self.wb = load_workbook(template)
        self.tables = {t.displayName: (ws, t) for ws in self.wb.worksheets for t in ws.tables.values()}

    def write_table(self, table_name: str, df: pd.DataFrame, formats: dict[str, str] | None = None) -> None:
        """Remplace en-têtes et lignes du tableau nommé par `df`, puis redimensionne sa plage.

        Le nombre de colonnes doit être celui du template : sinon le template et le code ont divergé.
        """
        if table_name not in self.tables:
            raise KeyError(f"tableau {table_name!r} absent du template {self.template.name}")
        ws, table = self.tables[table_name]
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        if len(df.columns) != max_col - min_col + 1:
            raise ValueError(f"{table_name} : {len(df.columns)} colonnes, le template en attend "
                             f"{max_col - min_col + 1}")

        for row in ws.iter_rows(min_row=min_row + 1, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row:
                cell.value = None

        formats = formats or {}
        for j, name in enumerate(df.columns):
            ws.cell(min_row, min_col + j, str(name))
        for i, values in enumerate(df.itertuples(index=False), start=1):
            for j, (name, v) in enumerate(zip(df.columns, values)):
                cell = ws.cell(min_row + i, min_col + j, _cell_value(v))
                if name in formats:
                    cell.number_format = formats[name]
                elif isinstance(cell.value, dt.datetime):
                    cell.number_format = "dd/mm/yyyy"

        last_row = min_row + max(len(df), 1)
        table.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{last_row}"
        if table.autoFilter is not None:
            table.autoFilter.ref = table.ref
        table.tableColumns = [TableColumn(id=j + 1, name=str(n)) for j, n in enumerate(df.columns)]

    def save(self) -> None:
        self.wb.calculation.fullCalcOnLoad = True  # Excel recalcule les formules de l'outil à l'ouverture
        self.output.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(self.output)
