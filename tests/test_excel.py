import pandas as pd
import pytest
from openpyxl import load_workbook

from pharma.excel import sheets
from pharma.excel.writer import TemplateWriter
from pharma.quality.reconciliation import summarize
from pharma.stats.descriptive import monthly_sales, volumes_and_shares, yearly_trend


def test_write_table_resizes_and_clears(real_settings, tmp_path):
    out = tmp_path / "out.xlsx"
    w = TemplateWriter(real_settings.template, out)
    w.write_table("tbl_synthese", pd.DataFrame({"Indicateur": list("abcd"), "Valeur": 1, "Commentaire": ""}))
    w.write_table("tbl_synthese", pd.DataFrame({"Indicateur": ["x"], "Valeur": [2], "Commentaire": ["y"]}))
    w.save()
    ws = load_workbook(out)["Synthèse"]
    assert ws.tables["tbl_synthese"].ref == "B4:D5"
    assert ws["B6"].value is None                        # anciennes lignes effacées


def test_column_mismatch_rejected(real_settings, tmp_path):
    w = TemplateWriter(real_settings.template, tmp_path / "out.xlsx")
    with pytest.raises(ValueError):
        w.write_table("tbl_synthese", pd.DataFrame({"a": [1]}))
    with pytest.raises(KeyError):
        w.write_table("tbl_inexistant", pd.DataFrame({"a": [1]}))


def test_full_workbook(real_settings, real_clean, real_reconciliation, real_stage, tmp_path):
    clean_df, report = real_clean
    monthly = monthly_sales(clean_df)
    data = {"clean": clean_df, "report": report.to_frame(), "reconciliation_summary": summarize(real_reconciliation),
            "volumes": volumes_and_shares(monthly), "yearly_trend": yearly_trend(monthly), **real_stage.outputs()}
    out = tmp_path / "Analyse_Pharma.xlsx"
    w = TemplateWriter(real_settings.template, out)
    sheets.write_summary(w, data, real_stage.last_complete_month)
    sheets.write_descriptive(w, data["volumes"], data["variability"], data["yearly_trend"])
    sheets.write_seasonality(w, data["seasonal_profile"], real_settings.atc_groups)
    sheets.write_forecast(w, data["forecast_next"], data["backtest_results"], data["external_effect"])
    sheets.write_data_quality(w, data["report"], data["reconciliation_summary"])
    sheets.write_sales(w, clean_df, real_settings.atc_groups)
    w.save()

    wb = load_workbook(out)
    for ws in wb.worksheets:
        for t in ws.tables.values():
            first_data_row = ws[t.ref.split(":")[0]].row + 1
            col = ws[t.ref.split(":")[0]].column
            assert ws.cell(first_data_row, col).value is not None, f"{t.displayName} vide"

    data_ws = wb["Data"]
    assert data_ws.tables["tbl_ventes"].ref == f"A1:M{1 + clean_df['date'][clean_df['source'] == 'daily'].nunique()}"
    assert data_ws["M2"].value.startswith("=SUMPRODUCT(B2:I2*(Outil!$C$9:$J$9")
    assert wb["Outil"]["C10"].value.startswith("=IF(C$9=")    # l'outil du template n'est pas écrasé
    assert wb.calculation.fullCalcOnLoad
